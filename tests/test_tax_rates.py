from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from ozon_app.aggregation import aggregate_calculations
from ozon_app.calculator import calculate_run
from ozon_app.database import Database
from ozon_app.excel_reader import parse_report
from ozon_app.exporter import export_calculation, export_run
from ozon_app.models import (
    AccrualRow,
    AdditionalIncomeRow,
    ParsedSource,
    Product,
    RealizationRow,
)
from ozon_app.tax_rates import (
    AppliedTaxRate,
    TaxRateError,
    TaxRatePeriod,
    TaxRateSchedule,
    format_applied_rates,
    merge_applied_rates,
)


def _schedule(*rows: tuple[date | None, float]) -> TaxRateSchedule:
    return TaxRateSchedule(TaxRatePeriod(valid_from, rate) for valid_from, rate in rows)


def _accrual(
    row_number: int,
    day: date,
    accrual_type: str,
    amount: float,
    *,
    article: str = "A",
    group: str = "Продажи",
    quantity: float = 0,
    price: float = 0,
) -> AccrualRow:
    return AccrualRow(
        source_name="Начисления.xlsx",
        sheet_name="Начисления",
        row_number=row_number,
        accrual_id=f"op-{row_number}",
        accrual_date=day,
        accrual_type=accrual_type,
        article=article,
        sku="1001" if article else "",
        product_name="Тестовый товар" if article else "",
        quantity=quantity,
        seller_price=price,
        amount=amount,
        service_group=group,
    )


def _accrual_source(rows: list[AccrualRow], name: str = "Начисления.xlsx") -> ParsedSource:
    dates = [row.accrual_date for row in rows if row.accrual_date]
    return ParsedSource(
        path=Path(name),
        file_hash=f"hash-{name}",
        report_type="ACCRUAL",
        sheet_name="Начисления",
        header_row=1,
        accrual_rows=rows,
        period_start=min(dates),
        period_end=max(dates),
    )


def _realization_source(
    name: str,
    start: date,
    end: date,
    rows: list[tuple[str, float, date | None]],
) -> ParsedSource:
    return ParsedSource(
        path=Path(name),
        file_hash=f"hash-{name}",
        report_type="REALIZATION",
        sheet_name="Детализация",
        header_row=10,
        realization_rows=[
            RealizationRow(
                source_name=name,
                sheet_name="Детализация",
                row_number=index + 13,
                raw_article="A",
                sku="",
                product_name="Тестовый товар",
                shipment=shipment,
                unit_price=amount,
                quantity=1,
                amount=amount,
                sale_date=sale_date,
            )
            for index, (shipment, amount, sale_date) in enumerate(rows)
        ],
        period_start=start,
        period_end=end,
    )


def _pdf_act(day: date, amount: float) -> ParsedSource:
    return ParsedSource(
        path=Path("Акт о премии.pdf"),
        file_hash="hash-act",
        report_type="ADDITIONAL_INCOME",
        sheet_name="PDF",
        header_row=0,
        additional_income_rows=[
            AdditionalIncomeRow(
                source_name="Акт о премии.pdf",
                page_number=1,
                document_number="1",
                income_date=day,
                income_type="Премия Ozon",
                amount=amount,
            )
        ],
        period_start=day,
        period_end=day,
    )


PRODUCTS = {"A": Product("A", "Тестовый товар", 60, 40)}


def _save(database: Database, calculation) -> int:
    return database.save_run(
        calculation,
        {str(source.path): source.path for source in calculation.source_files},
    )
MID_AUGUST = _schedule((None, 0.04), (date(2026, 8, 15), 0.06))


def _mid_month_calculation():
    """Август 2026, ставка 4% до 14.08 включительно и 6% с 15.08."""
    accruals = _accrual_source(
        [
            _accrual(2, date(2026, 8, 14), "Выручка", 1_000, quantity=1, price=1_000),
            _accrual(3, date(2026, 8, 15), "Выручка", 2_000, quantity=1, price=2_000),
            _accrual(4, date(2026, 8, 14), "Программы партнёров", 100),
            _accrual(5, date(2026, 8, 16), "Возврат выручки", -500, group="Возвраты", quantity=1),
            _accrual(6, date(2026, 8, 14), "Баллы за скидки", 50),
            _accrual(7, date(2026, 8, 20), "Баллы за скидки", 70),
            _accrual(8, date(2026, 8, 10), "Потеря по вине Ozon в логистике", 200, article="", group=""),
            _accrual(9, date(2026, 8, 20), "Премия Ozon", 300, article="", group=""),
            _accrual(10, date(2026, 8, 31), "Подписка Premium", -100, article="", group=""),
        ]
    )
    buyout_details = _realization_source(
        "BuyoutDetails_01.08.2026-15.08.2026.xlsx",
        date(2026, 8, 1),
        date(2026, 8, 15),
        [("S-1", 150, date(2026, 8, 14)), ("S-2", 250, date(2026, 8, 15))],
    )
    realization_without_dates = _realization_source(
        "RealizationReportCIS-1.xlsx",
        date(2026, 8, 1),
        date(2026, 8, 31),
        [("S-3", 300, None)],
    )
    return calculate_run(
        [accruals, buyout_details, realization_without_dates, _pdf_act(date(2026, 8, 25), 400)],
        PRODUCTS,
        MID_AUGUST,
    )


class TaxRateScheduleTests(unittest.TestCase):
    def test_rate_on_date_uses_latest_row_not_after_the_date(self) -> None:
        schedule = _schedule(
            (date(2026, 8, 1), 0.06),
            (None, 0.04),
            (date(2027, 1, 1), 0.07),
        )

        self.assertEqual(schedule.rate_on(date(2020, 1, 1)), 0.04)
        self.assertEqual(schedule.rate_on(date(2026, 7, 31)), 0.04)
        self.assertEqual(schedule.rate_on(date(2026, 8, 1)), 0.06)
        self.assertEqual(schedule.rate_on(date(2026, 12, 31)), 0.06)
        self.assertEqual(schedule.rate_on(date(2027, 1, 1)), 0.07)
        self.assertEqual(schedule.rate_on(date(2030, 5, 5)), 0.07)
        self.assertEqual([item.valid_from for item in schedule.periods], [None, date(2026, 8, 1), date(2027, 1, 1)])

    def test_validation_rules(self) -> None:
        invalid = {
            "пустой": [],
            "без строки с начала": [TaxRatePeriod(date(2026, 1, 1), 0.04)],
            "две строки с начала": [TaxRatePeriod(None, 0.04), TaxRatePeriod(None, 0.06)],
            "повтор даты": [
                TaxRatePeriod(None, 0.04),
                TaxRatePeriod(date(2026, 8, 1), 0.06),
                TaxRatePeriod(date(2026, 8, 1), 0.07),
            ],
            "ставка больше 100%": [TaxRatePeriod(None, 1.01)],
            "отрицательная ставка": [TaxRatePeriod(None, -0.01)],
            "не число": [TaxRatePeriod(None, float("nan"))],
        }
        for name, periods in invalid.items():
            with self.subTest(name), self.assertRaises(TaxRateError):
                TaxRateSchedule(periods)
        self.assertEqual(TaxRateSchedule([TaxRatePeriod(None, 0)]).rate_on(date.today()), 0)
        self.assertEqual(TaxRateSchedule([TaxRatePeriod(None, 1)]).rate_on(date.today()), 1)

    def test_changes_segments_and_comparison(self) -> None:
        schedule = _schedule((None, 0.04), (date(2026, 8, 15), 0.06), (date(2026, 9, 1), 0.06))

        self.assertEqual(
            schedule.rate_changes(date(2026, 8, 1), date(2026, 8, 31)),
            [(date(2026, 8, 15), 0.04, 0.06)],
        )
        self.assertEqual(schedule.rate_changes(date(2026, 8, 15), date(2026, 8, 31)), [])
        self.assertEqual(
            schedule.applied_rates(date(2026, 8, 1), date(2026, 9, 30)),
            [
                AppliedTaxRate(date(2026, 8, 1), date(2026, 8, 14), 0.04),
                AppliedTaxRate(date(2026, 8, 15), date(2026, 9, 30), 0.06),
            ],
        )
        self.assertTrue(
            schedule.same_rates_between(TaxRateSchedule.single(0.04), date(2026, 7, 1), date(2026, 8, 14))
        )
        self.assertFalse(
            schedule.same_rates_between(TaxRateSchedule.single(0.04), date(2026, 7, 1), date(2026, 8, 15))
        )
        self.assertFalse(schedule.same_rates_between(TaxRateSchedule.single(0.04), None, None))
        self.assertEqual(
            format_applied_rates(
                merge_applied_rates(
                    [
                        AppliedTaxRate(date(2026, 8, 1), date(2026, 8, 31), 0.06),
                        AppliedTaxRate(date(2026, 7, 1), date(2026, 7, 31), 0.04),
                        AppliedTaxRate(date(2026, 6, 1), date(2026, 6, 30), 0.04),
                    ]
                )
            ),
            "4 % — 01.06.2026–31.07.2026; 6 % — 01.08.2026–31.08.2026",
        )


class DatedTaxCalculationTests(unittest.TestCase):
    def test_rate_change_from_the_15th_taxes_each_row_by_its_own_date(self) -> None:
        calculation = _mid_month_calculation()
        product = calculation.products[0]

        # Выручка 1000×4% + 2000×6%, партнёры 100×4%, возврат −500×6%,
        # выкупы с датой 150×4% + 250×6%, выкуп без даты 300×6% (на 31.08).
        self.assertAlmostEqual(product.tax(calculation.tax_rate), 40 + 120 + 4 - 30 + 6 + 15 + 18)
        # Баллы (50 и 70) в базу не входят.
        self.assertAlmostEqual(product.taxable_income, 1_000 + 2_000 + 100 - 500 + 150 + 250 + 300)
        # Нераспределенные: 200×4% + 300×6% + PDF-акт 400×6%; подписка −100 не облагается.
        self.assertAlmostEqual(calculation.taxable_unallocated_income, 900)
        self.assertAlmostEqual(calculation.unallocated_income_tax, 8 + 18 + 24)
        self.assertAlmostEqual(calculation.totals()["tax"], 173 + 50)
        self.assertEqual(calculation.tax_rate, 0.06)
        self.assertEqual(
            calculation.applied_tax_rates,
            [
                AppliedTaxRate(date(2026, 8, 1), date(2026, 8, 14), 0.04),
                AppliedTaxRate(date(2026, 8, 15), date(2026, 8, 31), 0.06),
            ],
        )
        [warning] = calculation.tax_rate_warnings
        self.assertIn("RealizationReportCIS-1.xlsx", warning)
        self.assertIn("01.08.2026–31.08.2026", warning)
        self.assertIn("с 15.08.2026: 4 % → 6 %", warning)
        self.assertIn("по ставке на 31.08.2026 — 6 %", warning)

    def test_buyout_details_within_one_rate_period_needs_no_warning(self) -> None:
        accruals = _accrual_source(
            [_accrual(2, date(2026, 8, 20), "Выручка", 1_000, quantity=1, price=1_000)]
        )
        second_half = _realization_source(
            "RealizationReportCIS-2.xlsx",
            date(2026, 8, 16),
            date(2026, 8, 31),
            [("S-9", 500, None)],
        )

        calculation = calculate_run([accruals, second_half], PRODUCTS, MID_AUGUST)

        self.assertEqual(calculation.tax_rate_warnings, [])
        self.assertAlmostEqual(calculation.products[0].tax(calculation.tax_rate), 90)

    def test_single_rate_matches_previous_base_times_rate(self) -> None:
        calculation = _mid_month_calculation()
        single = calculate_run(
            [source for source in calculation.source_files],
            PRODUCTS,
            0.04,
        )
        product = single.products[0]

        self.assertAlmostEqual(product.tax(single.tax_rate), product.taxable_income * 0.04, places=9)
        self.assertAlmostEqual(
            single.unallocated_income_tax,
            single.taxable_unallocated_income * 0.04,
            places=9,
        )
        self.assertEqual(single.tax_rate_warnings, [])


class BuyoutDetailsDateTests(unittest.TestCase):
    def _write_buyout_report(self, path: Path, *, with_dates: bool) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Детализация"
        sheet["B2"] = "Детализация по выкупленным товарам за период с 01.08.2026 по 31.08.2026"
        sheet["B5"] = "Покупатель:"
        sheet["G5"] = "Продавец:"
        sheet["B6"] = "ООО «Тестовый покупатель»"
        sheet["G6"] = "Тестовый продавец"
        headers = {
            2: "№ п/п",
            3: "Товар",
            4: "Код товара продавца",
            5: "Код товара OZON",
            6: "Номер отправления",
            7: "Отчёт о выкупленных товарах" if with_dates else "Данные для расчета цены",
            9: "Реализовано",
        }
        for column, title in headers.items():
            sheet.cell(10, column, title)
        subheaders = {
            7: "Номер" if with_dates else "Исходные стоимостные данные",
            8: "Дата" if with_dates else "Дисконт по категории, %",
            9: "Цена реализации с НДС, руб.",
            10: "Ставка НДС, %",
            11: "Кол-во",
            12: "Итого к начислению, руб.",
        }
        for column, title in subheaders.items():
            sheet.cell(11, column, title)
        for column in range(2, 13):
            sheet.cell(12, column, column - 1)
        rows = [
            (1, "Тестовый товар", "A", 1001, "S-1", "100", "14.08.2026", 150, 1, 150),
            (2, "Тестовый товар", "A", 1001, "S-2", "101", "15.08.2026", 125, 2, 250),
        ]
        for offset, (number, name, article, sku, shipment, doc, day, price, qty, total) in enumerate(rows):
            row = 13 + offset
            values = [number, name, article, sku, shipment, doc if with_dates else 200, day if with_dates else 21, price, "Без НДС", qty, total]
            for column, value in enumerate(values, start=2):
                sheet.cell(row, column, value)
        sheet.cell(15, 2, "Итого с НДС:")
        sheet.cell(15, 12, 400)
        workbook.save(path)
        workbook.close()

    def test_parser_reads_row_dates_only_when_the_date_column_exists(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            with_dates = Path(directory) / "BuyoutDetails_01.08.2026-31.08.2026.xlsx"
            without_dates = Path(directory) / "RealizationReportCIS-2.xlsx"
            self._write_buyout_report(with_dates, with_dates=True)
            self._write_buyout_report(without_dates, with_dates=False)

            dated = parse_report(with_dates)
            undated = parse_report(without_dates)

        self.assertEqual(
            [row.sale_date for row in dated.realization_rows],
            [date(2026, 8, 14), date(2026, 8, 15)],
        )
        self.assertEqual([row.sale_date for row in undated.realization_rows], [None, None])
        self.assertEqual((dated.period_start, dated.period_end), (date(2026, 8, 1), date(2026, 8, 31)))

        accruals = _accrual_source(
            [_accrual(2, date(2026, 8, 20), "Выручка", 1_000, quantity=1, price=1_000)]
        )
        calculation = calculate_run([accruals, dated], PRODUCTS, MID_AUGUST)
        self.assertAlmostEqual(calculation.products[0].tax(calculation.tax_rate), 60 + 6 + 15)
        self.assertEqual(calculation.tax_rate_warnings, [])


class CombinedOverviewTaxTests(unittest.TestCase):
    def test_months_with_different_rates_sum_saved_taxes(self) -> None:
        schedule = _schedule((None, 0.04), (date(2026, 8, 1), 0.06))
        july = calculate_run(
            [
                _accrual_source(
                    [
                        _accrual(2, date(2026, 7, 1), "Выручка", 1_000, quantity=1, price=1_000),
                        _accrual(3, date(2026, 7, 31), "Потеря по вине Ozon в логистике", 1_000, article="", group=""),
                    ],
                    "Июль.xlsx",
                )
            ],
            PRODUCTS,
            schedule,
        )
        august = calculate_run(
            [
                _accrual_source(
                    [
                        _accrual(2, date(2026, 8, 1), "Выручка", 1_000, quantity=1, price=1_000),
                        _accrual(3, date(2026, 8, 31), "Подписка Premium", -10, article="", group=""),
                    ],
                    "Август.xlsx",
                )
            ],
            PRODUCTS,
            schedule,
        )
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "app.sqlite3")
            july_id = _save(database, july)
            august_id = _save(database, august)
            saved = [database.load_calculation(august_id), database.load_calculation(july_id)]

            combined = aggregate_calculations(saved)

            # Июль: 1000×4% + компенсация 1000×4%; август: 1000×6%.
            self.assertAlmostEqual(combined.totals()["tax"], 140)
            self.assertAlmostEqual(combined.unallocated_income_tax, 40)
            # Прежняя средневзвешенная ставка давала 100 + 1000 × 5% = 150.
            self.assertNotAlmostEqual(combined.totals()["tax"], 150)
            self.assertEqual(combined.tax_rate, 0.06)
            self.assertEqual(
                format_applied_rates(combined.applied_tax_rates),
                "4 % — 01.07.2026–31.07.2026; 6 % — 01.08.2026–31.08.2026",
            )
            self.assertAlmostEqual(
                combined.report_net_profit,
                sum(item.report_net_profit for item in saved),
            )

            destination = Path(directory) / "combined.xlsx"
            export_calculation(combined, destination)
            workbook = load_workbook(destination)
            try:
                sheet = workbook["КонсОтчет"]
                self.assertEqual(
                    sheet["C5"].value,
                    "Ставки налога: 4 % — 01.07.2026–31.07.2026; 6 % — 01.08.2026–31.08.2026",
                )
                self.assertEqual(sheet["P3"].value, "Ставка налога для сценария (на 31.08.2026)")
                self.assertEqual(sheet["P4"].value, 0.06)
                self.assertAlmostEqual(sheet["O8"].value, 100)
                self.assertAlmostEqual(sheet["J4"].value, 40)
            finally:
                workbook.close()


class TaxExportTests(unittest.TestCase):
    def test_export_writes_tax_values_applied_rates_and_scenario_rate_caption(self) -> None:
        calculation = _mid_month_calculation()
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "app.sqlite3")
            run_id = _save(database, calculation)
            destination = Path(directory) / "август.xlsx"

            export_run(database, run_id, destination)

            workbook = load_workbook(destination)
            try:
                sheet = workbook["КонсОтчет"]
                # Значения, а не формулы «база × одна ставка».
                self.assertIsInstance(sheet["O8"].value, (int, float))
                self.assertAlmostEqual(sheet["O8"].value, 173)
                self.assertIsInstance(sheet["J4"].value, (int, float))
                self.assertAlmostEqual(sheet["J4"].value, 50)
                self.assertEqual(sheet["P4"].value, 0.06)
                self.assertEqual(sheet["P4"].number_format, "0.00%")
                self.assertEqual(sheet["P3"].value, "Ставка налога для сценария (на 31.08.2026)")
                self.assertEqual(
                    sheet["C5"].value,
                    "Ставки налога: 4 % — 01.08.2026–14.08.2026; 6 % — 15.08.2026–31.08.2026",
                )
                self.assertEqual(sheet["AV8"].value, '=IF(AU8="","",AU8*$P$4)')
                self.assertEqual(sheet["M3"].value, "Чистая прибыль от деятельности")
                breakdown = {
                    row[0]: row[2]
                    for row in workbook["Разбивка"].iter_rows(values_only=True)
                    if row and row[0]
                }
                self.assertAlmostEqual(breakdown["Налог с нераспределенных доходов"], 50)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
