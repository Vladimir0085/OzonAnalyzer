"""Interface changes of 0.5.31 on synthetic data only."""

from __future__ import annotations

import os
import tempfile
import tkinter as tk
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from ozon_app.calculator import calculate_scenario
from ozon_app.cost_catalog_tab import (
    SHOW_SKIPPED_ROWS_TEXT,
    catalog_positions_without_full_cost,
    cost_catalog_warning_text,
    rows_without_full_cost,
)
from ozon_app.display_labels import NO_COST_TEXT, NO_SALES_TEXT, profitability_label
from ozon_app.display_modes import (
    EXPANDED_DETACH_TEXT,
    EXPANDED_FULLSCREEN_TEXT,
    NORMAL_DETACH_TEXT,
    NORMAL_FULLSCREEN_TEXT,
    LaptopFriendlyOZPriceAnalyzerApp,
)
from ozon_app.exporter import export_calculation
from ozon_app.help_content import (
    CATEGORY_CARD_HELP,
    OVERVIEW_CARD_HELP,
    OVERVIEW_HELP_CONTENT,
    REPORTS_HELP_CONTENT,
    REVENUE_SHARE_CARD_HELP,
    card_tooltip_text,
    help_plain_text,
)
from ozon_app.models import Product, ProductResult, RunCalculation, RunSummary
from ozon_app.report_totals import (
    POINTS_CARD_TITLE,
    ReportTotalsOZPriceAnalyzerApp,
    report_total_profitability_text,
)
from ozon_app.service import AppService
from ozon_app.ui import (
    _result_values,
    _scenario_values,
    filter_product_results,
    report_name_is_custom,
    run_full_display,
    run_header_display,
    run_tooltip_text,
)
from ozon_app.ui_containers import wrapped_positions
from tests.test_calculation_etalon import (
    AUGUST_ROWS,
    JULY_ROWS,
    PRODUCTS,
    TAX_SCHEDULE,
    _write_accrual_report,
)


PROFITABILITY_INDEX = 9  # «Доходность» in OVERVIEW_COLUMN_SPECS
SCENARIO_PROFITABILITY_INDEX = 8  # «Доходность» in SCENARIO_COLUMN_SPECS


def _run(run_id: int, name: str, start: str | None, end: str | None) -> RunSummary:
    return RunSummary(
        id=run_id,
        created_at="2026-09-01T10:00:00",
        period_start=start,
        period_end=end,
        source_count=1,
        units=0,
        revenue=0,
        net_profit=0,
        unallocated_total=0,
        status="OK",
        report_name=name,
    )


def _result(article: str, *, units: float, total_cost: float, revenue: float, result: float) -> ProductResult:
    return ProductResult(
        article=article,
        name=f"Товар {article}",
        material_cost=total_cost,
        labor_cost=0,
        units=units,
        revenue_no_points=revenue,
        financial_result=result,
        tax_override=0,
    )


class WrappingToolbarTests(unittest.TestCase):
    def test_right_aligned_rows_end_at_the_toolbar_edge(self) -> None:
        sizes = [(300, 30), (200, 30), (150, 30), (130, 30)]
        positions, height = wrapped_positions(sizes, 1000, align_right=True)
        self.assertEqual(height, 30)
        self.assertEqual({y for _x, y in positions}, {0})
        last_x, _ = positions[-1]
        self.assertEqual(last_x + sizes[-1][0], 1000)
        # Order is kept: «О программе» stays right of «Экспорт в Excel».
        self.assertEqual([x for x, _y in positions], sorted(x for x, _y in positions))

    def test_right_aligned_groups_wrap_without_clipping(self) -> None:
        sizes = [(300, 30), (200, 30), (150, 30), (130, 30)]
        for width in (320, 460, 700):
            with self.subTest(width=width):
                positions, height = wrapped_positions(sizes, width, align_right=True)
                self.assertGreater(height, 30)
                for (x, y), (group_width, group_height) in zip(positions, sizes):
                    self.assertGreaterEqual(x, 0)
                    self.assertLessEqual(x + group_width, width)
                    self.assertLessEqual(y + group_height, height)

    def test_left_alignment_is_unchanged_by_default(self) -> None:
        positions, _height = wrapped_positions([(100, 20), (100, 20)], 500)
        self.assertEqual(positions, [(0, 0), (108, 0)])


class ReportListLabelTests(unittest.TestCase):
    def test_generated_name_is_shown_as_compact_period(self) -> None:
        run = _run(7, "Отчет Ozon за 01.08.2026–31.08.2026", "2026-08-01", "2026-08-31")
        self.assertFalse(report_name_is_custom(run))
        self.assertEqual(run_header_display(4, run), "№4 · 01.08–31.08.2026")

    def test_legacy_generated_name_is_not_custom(self) -> None:
        run = _run(7, "Отчет Ozon #7", "2026-08-01", "2026-08-31")
        self.assertFalse(report_name_is_custom(run))
        self.assertEqual(run_header_display(4, run), "№4 · 01.08–31.08.2026")

    def test_renamed_report_keeps_its_name(self) -> None:
        run = _run(7, "Август — проверено", "2026-08-01", "2026-08-31")
        self.assertTrue(report_name_is_custom(run))
        self.assertEqual(run_header_display(4, run), "№4 · Август — проверено")

    def test_period_across_years_and_missing_period(self) -> None:
        crossing = _run(3, "Отчет Ozon за 16.12.2025–15.01.2026", "2025-12-16", "2026-01-15")
        self.assertEqual(run_header_display(1, crossing), "№1 · 16.12.2025–15.01.2026")
        missing = _run(5, "Отчет Ozon без периода", None, None)
        self.assertEqual(run_header_display(2, missing), "№2 · Период не определен")

    def test_tooltip_and_comparison_list_keep_the_full_name(self) -> None:
        run = _run(7, "Отчет Ozon за 01.08.2026–31.08.2026", "2026-08-01", "2026-08-31")
        self.assertEqual(
            run_tooltip_text(4, run),
            "Отчет №4: Отчет Ozon за 01.08.2026–31.08.2026\nПериод: 01.08.2026–31.08.2026",
        )
        self.assertEqual(
            run_full_display(4, run),
            "№4 · Отчет Ozon за 01.08.2026–31.08.2026 · 01.08.2026–31.08.2026",
        )


class ProfitabilityLabelTests(unittest.TestCase):
    def setUp(self) -> None:
        self.normal = _result("A", units=2, total_cost=100, revenue=500, result=400)
        self.no_sales = _result("B", units=0, total_cost=100, revenue=0, result=-35)
        self.no_cost = _result("C", units=1, total_cost=0, revenue=800, result=730)
        self.returned = _result("D", units=-1, total_cost=200, revenue=-450, result=-495)

    def test_wb_rule(self) -> None:
        self.assertEqual(profitability_label(0, 100), NO_SALES_TEXT)
        self.assertEqual(profitability_label(-1, -200), NO_SALES_TEXT)
        self.assertEqual(profitability_label(1, 0), NO_COST_TEXT)
        self.assertIsNone(profitability_label(1, 100))

    def test_overview_table_labels_replace_zero_percent_only_in_text(self) -> None:
        self.assertEqual(_result_values(self.normal, 0.04)[PROFITABILITY_INDEX], "100.00%")
        self.assertEqual(_result_values(self.no_sales, 0.04)[PROFITABILITY_INDEX], NO_SALES_TEXT)
        self.assertEqual(_result_values(self.no_cost, 0.04)[PROFITABILITY_INDEX], NO_COST_TEXT)
        self.assertEqual(_result_values(self.returned, 0.04)[PROFITABILITY_INDEX], NO_SALES_TEXT)
        # The numbers behind the labels are untouched.
        self.assertEqual(self.no_sales.profitability(0.04), 0.0)
        self.assertEqual(self.no_cost.profitability(0.04), 0.0)
        self.assertAlmostEqual(self.returned.profitability(0.04), 1.475)

    def test_scenario_table_labels(self) -> None:
        values = {
            result.article: _scenario_values(calculate_scenario(result, 0.04))[SCENARIO_PROFITABILITY_INDEX]
            for result in (self.normal, self.no_sales, self.no_cost, self.returned)
        }
        self.assertEqual(values["B"], NO_SALES_TEXT)
        self.assertEqual(values["C"], NO_COST_TEXT)
        self.assertEqual(values["D"], NO_SALES_TEXT)
        self.assertTrue(values["A"].endswith("%"))

    def test_sorting_by_profitability_uses_the_same_numbers(self) -> None:
        rows = [self.normal, self.no_sales, self.no_cost, self.returned]
        for descending in (False, True):
            with self.subTest(descending=descending):
                expected = sorted(
                    rows,
                    key=lambda row: (row.profitability(0.04), row.article.casefold()),
                    reverse=descending,
                )
                actual = filter_product_results(
                    rows, 0.04, sort_metric="Доходность", descending=descending
                )
                self.assertEqual([row.article for row in actual], [row.article for row in expected])

    def test_report_total_profitability_card(self) -> None:
        def calculation(products) -> RunCalculation:
            return RunCalculation(
                run_id=1,
                period_start=date(2026, 8, 1),
                period_end=date(2026, 8, 31),
                tax_rate=0.04,
                products=products,
                unallocated_total=100,
                unallocated={},
                accrual_stats={},
            )

        self.assertEqual(report_total_profitability_text(calculation([self.no_sales])), NO_SALES_TEXT)
        self.assertEqual(report_total_profitability_text(calculation([self.no_cost])), NO_COST_TEXT)
        self.assertEqual(report_total_profitability_text(calculation([self.normal])), "150.00%")


class ExportLabelTests(unittest.TestCase):
    def test_profitability_cells_show_labels_and_keep_the_formula(self) -> None:
        calculation = RunCalculation(
            run_id=1,
            period_start=date(2026, 8, 1),
            period_end=date(2026, 8, 31),
            tax_rate=0.04,
            products=[
                _result("A", units=2, total_cost=100, revenue=500, result=400),
                _result("B", units=0, total_cost=100, revenue=0, result=-35),
            ],
            unallocated_total=0,
            unallocated={},
            accrual_stats={},
        )
        with tempfile.TemporaryDirectory() as directory:
            with patch.dict(os.environ, {"OZON_APP_DATA": directory}):
                path = export_calculation(calculation, Path(directory) / "report.xlsx")
            workbook = load_workbook(path)
            sheet = workbook["КонсОтчет"]
            self.assertEqual(
                sheet["I8"].value,
                '=IF(Q8<=0,"Нет продаж",IF(C8<=0,"Нет себестоимости",IFERROR(J8/C8,0)))',
            )
            self.assertEqual(
                sheet["AO9"].value,
                '=IF(Q9<=0,"Нет продаж",IF(C9<=0,"Нет себестоимости",'
                'IF(OR(Q9=0,AN9=""),"",IFERROR(AY9/C9,0))))',
            )
            self.assertEqual(sheet["Q9"].value, 0)
            workbook.close()


class HelpAndCardTextTests(unittest.TestCase):
    def test_every_card_tooltip_formula_is_printed_in_help(self) -> None:
        text = help_plain_text(OVERVIEW_HELP_CONTENT)
        for group in (OVERVIEW_CARD_HELP, REVENUE_SHARE_CARD_HELP, CATEGORY_CARD_HELP):
            for key, card_help in group.items():
                with self.subTest(card=key):
                    formulas, _note = card_help
                    self.assertTrue(formulas)
                    for formula in formulas:
                        self.assertIn(formula, text)
                        self.assertIn(formula, card_tooltip_text(card_help))

    def test_help_describes_labels_points_income_and_short_report_names(self) -> None:
        text = help_plain_text(OVERVIEW_HELP_CONTENT)
        self.assertNotIn("показывает 0,00%", text)
        self.assertIn("«Нет продаж»", text)
        self.assertIn("«Нет себестоимости»", text)
        self.assertIn("не меняет порядок строк", text)
        self.assertIn("Баллы Ozon (доход)", text)
        self.assertIn("доход продавца, а не расход", text)
        self.assertIn("«№4 · 01.08–31.08.2026»", text)
        self.assertIn("Наведите указатель на любую карточку", text)
        self.assertIn("Показать пропущенные строки", help_plain_text(REPORTS_HELP_CONTENT))

    def test_points_card_title(self) -> None:
        self.assertEqual(POINTS_CARD_TITLE, "Баллы Ozon (доход)")


class CostCatalogWarningTests(unittest.TestCase):
    def test_zero_cost_positions_are_listed(self) -> None:
        products = [Product("A", "Кружка", 600, 400), Product("C", "Набор", 0, 0)]
        self.assertEqual(
            catalog_positions_without_full_cost(products),
            ["Артикул C «Набор»: полная себестоимость 0 ₽"],
        )
        self.assertEqual(catalog_positions_without_full_cost(products[:1]), [])

    def test_rows_without_full_cost_are_found_in_rejected_xlsx(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Справочник.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Артикул", "Наименование", "Полная себестоимость, руб.", "Трудозатраты, руб."])
            sheet.append(["A", "Кружка", 1000, 400])
            sheet.append(["B", "Ваза", None, 100])
            sheet.append([None, None, None, None])
            sheet.append(["C", "Набор", "", None])
            workbook.save(path)
            workbook.close()
            self.assertEqual(
                rows_without_full_cost(path),
                [
                    "Строка 3, артикул B: не заполнена полная себестоимость",
                    "Строка 5, артикул C: не заполнена полная себестоимость",
                ],
            )

    def test_warning_text_only_when_rows_are_missing(self) -> None:
        self.assertEqual(cost_catalog_warning_text([], []), "")
        text = cost_catalog_warning_text(["x"], ["y", "z"], "Справочник.xlsx")
        self.assertIn("«Справочник.xlsx» не загружен", text)
        self.assertIn("строк без полной себестоимости — 2", text)
        self.assertIn("Позиций с полной себестоимостью 0 ₽: 1", text)


class InterfaceTkTests(unittest.TestCase):
    """Real Tk geometry; runs on Windows CI and under Xvfb, skipped without a display."""

    @classmethod
    def setUpClass(cls) -> None:
        try:
            probe = tk.Tk()
        except tk.TclError as exc:
            if "no display name" in str(exc) or "couldn't connect to display" in str(exc):
                raise unittest.SkipTest(f"Tk display is unavailable: {exc}") from exc
            raise
        probe.destroy()

    def setUp(self) -> None:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        root = Path(directory.name)
        env_patch = patch.dict(os.environ, {"OZON_APP_DATA": str(root / "app")})
        env_patch.start()
        self.addCleanup(env_patch.stop)
        service = AppService(root / "data")
        for product in PRODUCTS:
            service.db.save_product(product, source="Тест")
        service.db.replace_tax_rates(TAX_SCHEDULE)
        july = root / "Отчет по начислениям_01.07.2026-31.07.2026.xlsx"
        august = root / "Отчет по начислениям_01.08.2026-31.08.2026.xlsx"
        _write_accrual_report(july, JULY_ROWS)
        _write_accrual_report(august, AUGUST_ROWS)
        service.complete_import_batch(service.prepare_import([july, august]).sessions)
        self.july_id, self.august_id = [run.id for run in service.db.list_runs()]
        service.db.rename_run(self.july_id, "Июль — проверено")

        self.callback_errors: list[tuple[object, ...]] = []
        error_patch = patch(
            "tkinter.Tk.report_callback_exception",
            side_effect=lambda *error: self.callback_errors.append(error),
        )
        error_patch.start()
        self.addCleanup(error_patch.stop)
        initialize_tk = tk.Tk.__init__

        def initialize_at_test_dpi(tk_root, *args, **kwargs):
            initialize_tk(tk_root, *args, **kwargs)
            tk_root.tk.call("tk", "scaling", 96 / 72)

        with patch("tkinter.Tk.__init__", initialize_at_test_dpi), patch.object(
            LaptopFriendlyOZPriceAnalyzerApp, "_maximize_on_small_screen", lambda self: None
        ):
            self.app = ReportTotalsOZPriceAnalyzerApp(service)
        self.addCleanup(self.app.destroy)
        self.app.minsize(1, 1)
        self.app.maxsize(2200, 1400)
        self.app._base_tk_scaling = 96 / 72
        self.app._apply_ui_scale(1.0)

    def settle(self) -> None:
        for _ in range(4):
            self.app.update_idletasks()
            self.app.update()
        self.assertEqual(self.callback_errors, [])

    def assert_inside(self, child, parent) -> None:
        self.assertTrue(child.winfo_viewable(), str(child))
        self.assertGreater(child.winfo_width(), 1)
        self.assertGreaterEqual(child.winfo_rootx(), parent.winfo_rootx(), str(child))
        self.assertGreaterEqual(child.winfo_rooty(), parent.winfo_rooty(), str(child))
        self.assertLessEqual(
            child.winfo_rootx() + child.winfo_width(),
            parent.winfo_rootx() + parent.winfo_width(),
            str(child),
        )
        self.assertLessEqual(
            child.winfo_rooty() + child.winfo_height(),
            parent.winfo_rooty() + parent.winfo_height(),
            str(child),
        )

    def assert_groups_inside(self, toolbar, parent) -> None:
        self.assert_inside(toolbar, parent)
        for group in toolbar.groups:
            self.assert_inside(group, toolbar)
            for child in group.winfo_children():
                if child.winfo_manager():
                    self.assert_inside(child, group)

    def _button_texts(self, group) -> list[str]:
        return [
            str(child.cget("text"))
            for child in group.winfo_children()
            if child.winfo_class() == "TButton"
        ]

    def test_about_button_is_in_one_row_right_of_export(self) -> None:
        self.app.geometry("1540x920+0+0")
        self.settle()
        groups = self.app.header_actions.groups
        texts = [self._button_texts(group) for group in groups]
        self.assertEqual(texts[1:], [["Импортировать отчеты"], ["Экспорт в Excel"], ["О программе"]])
        self.assertEqual({group.winfo_y() for group in groups}, {0})
        self.assertGreater(groups[3].winfo_rootx(), groups[2].winfo_rootx() + groups[2].winfo_width())
        self.assert_groups_inside(self.app.header_actions, self.app.header_frame)
        self.assertLess(self.app.header_frame.winfo_height(), 110)

    def test_header_buttons_wrap_in_a_narrow_window(self) -> None:
        for width in (1000, 1180):
            with self.subTest(width=width):
                self.app.geometry(f"{width}x720+0+0")
                self.settle()
                self.assert_groups_inside(self.app.header_actions, self.app.header_frame)
                if width == 1000:
                    rows = {group.winfo_y() for group in self.app.header_actions.groups}
                    self.assertGreater(len(rows), 1)

    def test_report_list_shows_compact_labels_with_full_tooltip(self) -> None:
        self.settle()
        self.assertEqual(
            tuple(self.app.run_combo.cget("values")),
            ("№1 · Июль — проверено", "№2 · 02.08–31.08.2026"),
        )
        self.assertEqual(self.app.run_var.get(), "№2 · 02.08–31.08.2026")
        self.assertIn("Отчет Ozon за 02.08.2026–31.08.2026", self.app._run_combo_tooltip_text())
        self.app.run_var.set("№1 · Июль — проверено")
        self.app._on_run_selected()
        self.assertEqual(self.app.current_run_id, self.july_id)
        self.assertEqual(
            self.app._run_combo_tooltip_text(),
            "Отчет №1: Июль — проверено\nПериод: 03.07.2026–30.07.2026",
        )
        # The comparison lists keep the full names.
        self.assertIn(
            "№1 · Июль — проверено · 03.07.2026–30.07.2026",
            tuple(self.app.compare_first_combo.cget("values")),
        )

    def test_scenario_rows_wrap_so_reset_and_direction_stay_visible(self) -> None:
        self.app.notebook.select(self.app.scenario_tab)
        # Two reports make the counter as long as it gets in practice.
        self.app.overview_run_ids = {self.july_id, self.august_id}
        self.app.overview_selection_explicit = True
        self.app._refresh_active_report_views()
        self.assertIn("цены только для 1 отчета", self.app.scenario_count_var.get())
        for width, height in ((1540, 920), (1180, 720)):
            for scale in (1.0, 0.9):
                with self.subTest(width=width, height=height, scale=scale):
                    self.app._apply_ui_scale(scale)
                    self.app.geometry(f"{width}x{height}+0+0")
                    self.settle()
                    tab = self.app.scenario_tab
                    self.assert_groups_inside(self.app.scenario_filters, tab)
                    self.assert_groups_inside(self.app.scenario_price_controls, tab)
                    self.assert_inside(self.app.scenario_reset_button, tab)
                    self.assert_inside(self.app.scenario_direction_combo, tab)
                    self.assertGreaterEqual(
                        self.app.scenario_tree.winfo_rooty(),
                        self.app.scenario_filters.winfo_rooty() + self.app.scenario_filters.winfo_height(),
                    )

    def test_expanded_tables_keep_buttons_in_a_row_above_the_headings(self) -> None:
        self.app.geometry("1540x920+0+0")
        cases = (
            ("scenario", self.app.scenario_tab, self.app.scenario_tree, self.app._scenario_splitter),
            ("catalog", self.app.cost_catalog_tab, self.app.products_tree, self.app._catalog_splitter),
        )
        for key, tab, tree, splitter in cases:
            with self.subTest(table=key):
                self.app.notebook.select(tab)
                self.settle()
                mode = self.app._table_modes[key]
                self.assertFalse(hasattr(mode, "full_toolbar"))
                before = tree.winfo_height()
                mode.expand()
                self.settle()
                self.assertEqual(str(mode.fullscreen_button.cget("text")), EXPANDED_FULLSCREEN_TEXT)
                self.assertEqual(str(mode.detach_button.cget("text")), EXPANDED_DETACH_TEXT)
                self.assert_inside(mode.fullscreen_button, splitter.bar)
                self.assert_inside(mode.detach_button, splitter.bar)
                self.assertGreaterEqual(tree.winfo_rooty(), splitter.bar.winfo_rooty() + splitter.bar.winfo_height())
                self.assertGreater(tree.winfo_height(), before)
                if key == "scenario":
                    self.assertFalse(self.app.scenario_kpi_frame.winfo_viewable())
                mode.restore()
                self.settle()
                self.assertEqual(str(mode.fullscreen_button.cget("text")), NORMAL_FULLSCREEN_TEXT)
                self.assertEqual(str(mode.detach_button.cget("text")), NORMAL_DETACH_TEXT)
                if key == "scenario":
                    self.assertTrue(self.app.scenario_kpi_frame.winfo_viewable())

    def test_profitability_labels_in_tables_and_cards(self) -> None:
        self.app.run_var.set("№1 · Июль — проверено")
        self.app._on_run_selected()
        self.settle()
        tree = self.app.overview_tree
        self.assertEqual(tree.set("D", "profitability"), NO_SALES_TEXT)
        self.assertEqual(tree.set("C", "profitability"), NO_COST_TEXT)
        self.assertTrue(tree.set("A", "profitability").endswith("%"))
        self.assertEqual(self.app.scenario_tree.set("D", "profitability"), NO_SALES_TEXT)
        self.assertEqual(self.app.scenario_tree.set("C", "profitability"), NO_COST_TEXT)
        self.assertTrue(self.app.kpi_vars["profitability"].get().endswith("%"))
        self.assertTrue(self.app.kpi_vars["report_total_profitability"].get().endswith("%"))

    def test_every_overview_card_has_a_formula_tooltip(self) -> None:
        self.settle()
        cards = {
            **self.app.kpi_cards,
            **{f"share:{key}": card for key, card in self.app.revenue_share_kpi_cards.items()},
            **{f"category:{key}": card for key, card in self.app.category_kpi_cards.items()},
        }
        self.assertEqual(len(cards), 17)
        self.assertEqual(sorted(self.app.overview_card_tooltips), sorted(cards))
        help_text = help_plain_text(OVERVIEW_HELP_CONTENT)
        for key, tooltip in self.app.overview_card_tooltips.items():
            with self.subTest(card=key):
                self.assertIs(tooltip.widget, cards[key])
                first_line = tooltip.text().splitlines()[0]
                self.assertIn(first_line, help_text)
        tooltip = self.app.overview_card_tooltips["share:points"]
        tooltip.show(10, 10)
        self.settle()
        self.assertIsNotNone(tooltip.tipwindow)
        tooltip.hide()
        self.assertIsNone(tooltip.tipwindow)
        titles = [
            str(child.cget("text"))
            for child in self.app.revenue_share_kpi_cards["points"].winfo_children()
            if child.winfo_class() == "TLabel" and str(child.cget("text"))
        ]
        self.assertIn(POINTS_CARD_TITLE, titles)

    def test_catalog_warning_is_permanent_and_button_only_with_missing_rows(self) -> None:
        self.app.notebook.select(self.app.cost_catalog_tab)
        self.settle()
        frame = self.app.cost_catalog_warning_frame
        self.assertTrue(frame.winfo_viewable())
        self.assertEqual(str(self.app.cost_catalog_warning_button.cget("text")), SHOW_SKIPPED_ROWS_TEXT)
        self.assertIn("0 ₽: 1", self.app.cost_catalog_warning_var.get())
        self.assert_inside(self.app.cost_catalog_warning_button, self.app.cost_catalog_tab)

        self.app.db.save_product(Product("C", "Подарочный набор", 90, 10, category="Декор"), source="Тест")
        self.app.refresh_products()
        self.settle()
        self.assertFalse(frame.winfo_ismapped())
        self.assertEqual(self.app.cost_catalog_warning_var.get(), "")

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Справочник.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Артикул", "Наименование", "Полная себестоимость, руб.", "Трудозатраты, руб."])
            sheet.append(["A", "Кружка", 1000, 400])
            sheet.append(["N", "Новинка", None, None])
            workbook.save(path)
            workbook.close()
            products_before = self.app.db.list_products()
            with patch("ozon_app.ui.filedialog.askopenfilename", return_value=str(path)), patch(
                "ozon_app.ui.messagebox.showerror"
            ) as error:
                self.app.import_product_catalog()
            error.assert_called_once()
        self.settle()
        # The import rule is unchanged: the catalog stays as it was.
        self.assertEqual(self.app.db.list_products(), products_before)
        self.assertTrue(frame.winfo_viewable())
        self.assertIn("«Справочник.xlsx» не загружен", self.app.cost_catalog_warning_var.get())
        with patch("ozon_app.cost_catalog_tab.messagebox.showwarning") as warning:
            self.app.show_cost_catalog_warnings()
        message = warning.call_args.args[1]
        self.assertIn("Строка 3, артикул N: не заполнена полная себестоимость", message)


if __name__ == "__main__":
    unittest.main()
