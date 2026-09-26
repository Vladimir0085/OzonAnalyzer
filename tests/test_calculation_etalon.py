"""Эталон расчётов на синтетических отчётах.

Эталон ``fixtures/calculation_etalon.json`` снят на версии 0.5.30 до
интерфейсных изменений 0.5.31. Любое изменение интерфейса, подписей или
справки не должно менять ни одной копейки: выручку, продажи, налог, чистую
прибыль товаров, чистую прибыль от деятельности, доходность, доли от выручки,
сценарий цены и числовые ячейки экспорта XLSX.

Перезаписывать эталон можно только при осознанном изменении расчёта:
``python -m tests.test_calculation_etalon --write``.
"""

from __future__ import annotations

import json
import os
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from ozon_app.aggregation import aggregate_calculations
from ozon_app.calculator import calculate_scenario
from ozon_app.exporter import export_calculation
from ozon_app.models import Product, RunCalculation
from ozon_app.service import AppService
from ozon_app.tax_rates import TaxRatePeriod, TaxRateSchedule


FIXTURE = Path(__file__).resolve().parent / "fixtures" / "calculation_etalon.json"
MONEY_TOLERANCE = 1e-6
RATIO_TOLERANCE = 1e-9

HEADERS = [
    "ID начисления",
    "Дата начисления",
    "Группа услуг",
    "Тип начисления",
    "Артикул",
    "SKU",
    "Название товара",
    "Количество",
    "Цена продавца",
    "Сумма итого, руб.",
]

PRODUCTS = (
    Product("A", "Кружка", 600, 400, category="Посуда"),
    Product("B", "Ваза", 400, 100, category="Декор"),
    # Нулевая себестоимость: доходность отображается как «Нет себестоимости».
    Product("C", "Подарочный набор", 0, 0, category="Декор"),
    # Только расход без продаж: доходность отображается как «Нет продаж».
    Product("D", "Подставка", 250, 50),
    # В августе только возврат июльской продажи: продажи меньше нуля.
    Product("E", "Тарелка", 180, 20, category="Посуда"),
)
SKU = {"A": 1001, "B": 1002, "C": 1003, "D": 1004, "E": 1005}
NAMES = {product.article: product.name for product in PRODUCTS}
TAX_SCHEDULE = TaxRateSchedule(
    [TaxRatePeriod(None, 0.04), TaxRatePeriod(date(2026, 8, 15), 0.06)]
)
PLANNED_PRICES = {"A": 1_700.0, "B": 1_000.0}

TOTAL_KEYS = (
    "units",
    "revenue",
    "financial_result",
    "cost_sold",
    "taxable_income",
    "taxable_unallocated_income",
    "product_tax",
    "unallocated_income_tax",
    "tax",
    "net_profit",
    "unallocated",
    "report_net_profit",
)
PRODUCT_MONEY_KEYS = (
    "units",
    "revenue",
    "taxable_income",
    "financial_result",
    "cost_sold",
    "tax",
    "net_profit",
)
PRODUCT_RATIO_KEYS = (
    "profitability",
    "commission_share",
    "logistics_share",
    "points_share",
    "net_margin",
)
SCENARIO_KEYS = (
    "current_price",
    "planned_price",
    "price_change",
    "profitability",
    "planned_revenue",
    "commission_rate",
    "planned_commission",
    "planned_points",
    "taxable_base",
    "tax",
    "profit",
    "net_profit_per_unit",
    "net_profit_total",
)
RATIO_NAMES = set(PRODUCT_RATIO_KEYS) | {
    "report_profitability",
    "price_change",
    "commission_rate",
}
# Единственные ячейки экспорта, которые 0.5.31 может изменить: доходность и
# плановая доходность товара получают подпись «Нет продаж» / «Нет себестоимости».
_LABEL_CELL = re.compile(r"^КонсОтчет!(I|AO)(\d+)$")


def _sale(op: str, day: date, article: str, quantity: float, price: float) -> tuple[object, ...]:
    return (op, day, "Продажи", "Выручка", article, SKU[article], NAMES[article], quantity, price, quantity * price)


def _sale_part(op: str, day: date, accrual_type: str, article: str, amount: float) -> tuple[object, ...]:
    return (op, day, "Продажи", accrual_type, article, SKU[article], NAMES[article], 0, 0, amount)


def _charge(op: str, day: date, accrual_type: str, article: str, amount: float) -> tuple[object, ...]:
    return (op, day, "", accrual_type, article, SKU[article], NAMES[article], 0, 0, amount)


def _return(op: str, day: date, article: str, quantity: float, price: float) -> tuple[object, ...]:
    return (
        op, day, "Возвраты", "Возврат выручки", article, SKU[article], NAMES[article],
        quantity, price, -quantity * price,
    )


def _unallocated(op: str, day: date, accrual_type: str, amount: float) -> tuple[object, ...]:
    return (op, day, "", accrual_type, "", "", "", 0, 0, amount)


JULY_ROWS = (
    _sale("J-1", date(2026, 7, 3), "A", 2, 1_500),
    _sale_part("J-1", date(2026, 7, 3), "Баллы за скидки", "A", 240),
    _sale_part("J-1", date(2026, 7, 3), "Программы партнёров", "A", 60),
    _charge("J-1", date(2026, 7, 3), "Вознаграждение за продажу", "A", -540),
    _charge("J-2", date(2026, 7, 5), "Обработка отправления Drop-off", "A", -60),
    _charge("J-3", date(2026, 7, 5), "Доставка до места выдачи", "A", -90),
    _charge("J-4", date(2026, 7, 6), "Логистика", "A", -210),
    _sale("J-5", date(2026, 7, 10), "B", 3, 900),
    _charge("J-5", date(2026, 7, 10), "Вознаграждение за продажу", "B", -405),
    _charge("J-6", date(2026, 7, 11), "Эквайринг", "B", -40.5),
    _sale("J-7", date(2026, 7, 12), "C", 1, 800),
    _charge("J-8", date(2026, 7, 12), "Логистика", "C", -70),
    _sale("J-9", date(2026, 7, 20), "E", 1, 450),
    _charge("J-10", date(2026, 7, 25), "Логистика", "D", -35.25),
    _unallocated("J-11", date(2026, 7, 28), "Подписка Premium", -499),
    _unallocated("J-12", date(2026, 7, 30), "Премия Ozon", 1_200),
)
AUGUST_ROWS = (
    _sale("A-1", date(2026, 8, 2), "A", 1, 1_600),
    _sale_part("A-1", date(2026, 8, 2), "Баллы за скидки", "A", 150),
    _charge("A-1", date(2026, 8, 2), "Вознаграждение за продажу", "A", -240),
    _sale("A-2", date(2026, 8, 20), "A", 1, 1_650),
    _sale_part("A-2", date(2026, 8, 20), "Программы партнёров", "A", 50),
    _charge("A-2", date(2026, 8, 20), "Вознаграждение за продажу", "A", -247.5),
    _charge("A-2", date(2026, 8, 21), "Обратная логистика", "A", -80),
    _return("A-3", date(2026, 8, 18), "E", 1, 450),
    _charge("A-4", date(2026, 8, 18), "Обработка возвратов, отмен и невыкупов партнёрами", "E", -45),
    _sale("A-5", date(2026, 8, 10), "B", 2, 950),
    _charge("A-5", date(2026, 8, 10), "Вознаграждение за продажу", "B", -285),
    _charge("A-6", date(2026, 8, 12), "Звёздные товары", "B", -95),
    _charge("A-7", date(2026, 8, 22), "Упаковка товара партнёрами", "B", -30),
    _charge("A-8", date(2026, 8, 25), "Потеря по вине Ozon в логистике", "A", 500),
    _unallocated("A-9", date(2026, 8, 5), "Подписка Premium", -499),
    _unallocated("A-10", date(2026, 8, 16), "Премия Ozon", 800),
    _unallocated("A-11", date(2026, 8, 31), "Потеря по вине Ozon в логистике", 300),
)


def _write_accrual_report(path: Path, rows) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Начисления"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    workbook.close()


def _number(value) -> float | None:
    return None if value is None else float(value)


def _calculation_metrics(calculation: RunCalculation) -> dict[str, object]:
    totals = calculation.totals()
    cost_sold = totals["cost_sold"]
    rate = calculation.tax_rate
    products = {}
    for result in calculation.products:
        products[result.article] = {
            "units": result.units,
            "revenue": result.revenue_including_points,
            "taxable_income": result.taxable_income,
            "financial_result": result.financial_result,
            "cost_sold": result.cost_sold,
            "tax": result.tax(rate),
            "net_profit": result.net_profit(rate),
            "profitability": result.profitability(rate),
            "commission_share": result.commission_share(),
            "logistics_share": result.logistics_share(),
            "points_share": result.points_share(),
            "net_margin": result.net_margin(rate),
        }
    return {
        "tax_rate": rate,
        "totals": {key: totals[key] for key in TOTAL_KEYS},
        "report_profitability": calculation.report_net_profit / cost_sold if cost_sold else 0.0,
        "product_profitability": totals["net_profit"] / cost_sold if cost_sold else 0.0,
        "revenue_shares": calculation.revenue_shares(),
        "revenue_amounts": calculation.revenue_amounts(),
        "products": products,
    }


def _scenario_metrics(calculation: RunCalculation, prices: dict[str, float]) -> dict[str, object]:
    rows = {}
    for result in calculation.products:
        row = calculate_scenario(result, calculation.tax_rate, prices.get(result.article))
        rows[result.article] = {key: _number(getattr(row, key)) for key in SCENARIO_KEYS}
    return rows


def _export_cells(calculation: RunCalculation, destination: Path, prices=None) -> dict[str, object]:
    export_calculation(calculation, destination, prices)
    workbook = load_workbook(destination)
    try:
        cells: dict[str, object] = {}
        for sheet_name in ("КонсОтчет", "Разбивка"):
            for row in workbook[sheet_name].iter_rows():
                for cell in row:
                    if cell.value is not None:
                        cells[f"{sheet_name}!{cell.coordinate}"] = cell.value
        return cells
    finally:
        workbook.close()


def build_snapshot(root: Path) -> dict[str, object]:
    """Build every etalon value from synthetic reports through the real pipeline."""
    with patch.dict(os.environ, {"OZON_APP_DATA": str(root / "app")}):
        service = AppService(root / "data")
        for product in PRODUCTS:
            service.db.save_product(product, source="Эталон")
        service.db.replace_tax_rates(TAX_SCHEDULE)
        july = root / "Отчет по начислениям_01.07.2026-31.07.2026.xlsx"
        august = root / "Отчет по начислениям_01.08.2026-31.08.2026.xlsx"
        _write_accrual_report(july, JULY_ROWS)
        _write_accrual_report(august, AUGUST_ROWS)
        batch = service.prepare_import([july, august])
        service.complete_import_batch(batch.sessions)

        runs = service.db.list_runs()
        calculations = {
            f"{run.period_start}..{run.period_end}": service.db.load_calculation(run.id)
            for run in runs
        }
        august_run = runs[-1]
        for article, price in PLANNED_PRICES.items():
            service.db.save_planned_price(august_run.id, article, price)
        overview = aggregate_calculations(list(calculations.values()))

        snapshot: dict[str, object] = {
            "reports": {key: _calculation_metrics(value) for key, value in calculations.items()},
            "overview": _calculation_metrics(overview),
            "scenario": _scenario_metrics(
                service.db.load_calculation(august_run.id),
                service.db.planned_prices(august_run.id),
            ),
            "export": {},
        }
        exports_dir = root / "exports"
        for key, calculation in calculations.items():
            run_id = calculation.run_id
            snapshot["export"][key] = _export_cells(
                calculation,
                exports_dir / f"{key}.xlsx",
                service.db.planned_prices(run_id) if run_id is not None else {},
            )
        snapshot["export"]["overview"] = _export_cells(overview, exports_dir / "overview.xlsx")
        return snapshot


def _labelled_profitability_formula(original: str, row: str) -> str:
    return (
        f'=IF(Q{row}<=0,"Нет продаж",IF(C{row}<=0,"Нет себестоимости",'
        f"{original[1:]}))"
    )


class CalculationEtalonTests(unittest.TestCase):
    """До и после интерфейсных изменений все показатели совпадают до копейки."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.expected = json.loads(FIXTURE.read_text(encoding="utf-8"))
        directory = tempfile.TemporaryDirectory()
        cls.addClassCleanup(directory.cleanup)
        cls.actual = build_snapshot(Path(directory.name))

    def assert_values_match(self, expected, actual, path: str) -> None:
        if isinstance(expected, dict):
            self.assertIsInstance(actual, dict, path)
            self.assertEqual(sorted(actual), sorted(expected), path)
            for key in expected:
                self.assert_values_match(expected[key], actual[key], f"{path}.{key}")
            return
        if expected is None or isinstance(expected, str):
            self.assertEqual(actual, expected, path)
            return
        self.assertIsNotNone(actual, path)
        leaf = path.rsplit(".", 1)[-1]
        ratio = leaf in RATIO_NAMES or leaf.endswith("_share") or leaf in {"net_margin", "tax_rate"}
        tolerance = RATIO_TOLERANCE if ratio else MONEY_TOLERANCE
        self.assertLessEqual(abs(float(actual) - float(expected)), tolerance, path)

    def test_fixture_covers_every_display_edge_case(self) -> None:
        august = self.expected["reports"]["2026-08-02..2026-08-31"]["products"]
        july = self.expected["reports"]["2026-07-03..2026-07-30"]["products"]
        self.assertEqual(july["D"]["units"], 0)
        self.assertEqual(july["C"]["cost_sold"], 0)
        self.assertLess(august["E"]["units"], 0)
        self.assertNotEqual(self.expected["overview"]["totals"]["unallocated"], 0)

    def test_report_totals_taxes_profit_and_shares_match_etalon(self) -> None:
        for key in ("reports", "overview"):
            with self.subTest(section=key):
                self.assert_values_match(self.expected[key], self.actual[key], key)

    def test_price_scenario_matches_etalon(self) -> None:
        self.assert_values_match(self.expected["scenario"], self.actual["scenario"], "scenario")

    def test_export_cells_match_etalon(self) -> None:
        for workbook_key, expected_cells in self.expected["export"].items():
            actual_cells = self.actual["export"][workbook_key]
            with self.subTest(workbook=workbook_key):
                self.assertEqual(sorted(actual_cells), sorted(expected_cells))
                for coordinate, expected in expected_cells.items():
                    actual = actual_cells[coordinate]
                    label_cell = _LABEL_CELL.match(coordinate)
                    if (
                        label_cell
                        and int(label_cell.group(2)) >= 8
                        and isinstance(expected, str)
                        and actual != expected
                    ):
                        # Only the display wrapper may be added; the formula inside is unchanged.
                        self.assertEqual(
                            actual,
                            _labelled_profitability_formula(expected, label_cell.group(2)),
                            coordinate,
                        )
                        continue
                    if isinstance(expected, str):
                        self.assertEqual(actual, expected, coordinate)
                    else:
                        self.assertLessEqual(
                            abs(float(actual) - float(expected)), MONEY_TOLERANCE, coordinate
                        )


def _write_fixture() -> None:
    with tempfile.TemporaryDirectory() as directory:
        snapshot = build_snapshot(Path(directory))
    FIXTURE.parent.mkdir(parents=True, exist_ok=True)
    FIXTURE.write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=1, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(f"Эталон записан: {FIXTURE}")


if __name__ == "__main__":
    if "--write" in sys.argv:
        _write_fixture()
    else:
        unittest.main()
