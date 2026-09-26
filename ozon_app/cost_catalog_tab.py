from __future__ import annotations

import tkinter as tk
from pathlib import Path
from tkinter import messagebox, ttk
from typing import Iterable

from openpyxl import load_workbook

from .catalog_category_filters import CatalogAndCategoryOZPriceAnalyzerApp
from .costs import CostCatalogError, _find_catalog_sheet
from .display_modes import (
    _TableModeController,
    _button_with_text,
    resolve_ui_scale,
)
from .excel_reader import display_text, is_numeric, normalize_text
from .models import Product
from .resizable_layout import _GridSplitter, _hide_label_with_text


SHOW_SKIPPED_ROWS_TEXT = "Показать пропущенные строки"
_WARNING_LIST_LIMIT = 25


def catalog_positions_without_full_cost(products: Iterable[Product]) -> list[str]:
    """Catalog positions whose full cost is 0 ₽ (profitability shows «Нет себестоимости»)."""
    return [
        f"Артикул {product.article} «{product.name}»: полная себестоимость 0 ₽"
        for product in products
        if product.total_cost <= 0
    ]


def rows_without_full_cost(path: str | Path) -> list[str]:
    """Rows of a catalog XLSX without a full cost.

    Read-only diagnostics for the warning: the import rule itself is unchanged,
    a file with such rows is still rejected as a whole.
    """
    try:
        workbook = load_workbook(Path(path), read_only=True, data_only=True)
    except Exception:
        return []
    try:
        try:
            ws, header_row, columns = _find_catalog_sheet(workbook)
        except CostCatalogError:
            return []
        article_column = columns[normalize_text("Артикул")]
        name_column = columns[normalize_text("Наименование")]
        total_column = columns[normalize_text("Полная себестоимость, руб.")]
        labor_column = columns[normalize_text("Трудозатраты, руб.")]
        rows: list[str] = []
        for row_number in range(header_row + 1, int(ws.max_row or header_row) + 1):
            article = display_text(ws.cell(row_number, article_column).value)
            name = display_text(ws.cell(row_number, name_column).value)
            total_value = ws.cell(row_number, total_column).value
            labor_value = ws.cell(row_number, labor_column).value
            if not article and not name and total_value in (None, "") and labor_value in (None, ""):
                continue
            if not is_numeric(total_value):
                rows.append(
                    f"Строка {row_number}, артикул {article or 'не указан'}: "
                    "не заполнена полная себестоимость"
                )
        return rows
    finally:
        workbook.close()


def cost_catalog_warning_text(
    zero_cost: list[str],
    rejected_rows: list[str],
    rejected_source: str = "",
) -> str:
    """Text of the permanent warning above the catalog table ("" — no warning)."""
    parts: list[str] = []
    if rejected_rows:
        source = f" «{rejected_source}»" if rejected_source else ""
        parts.append(
            f"Последний XLSX{source} не загружен: строк без полной себестоимости — "
            f"{len(rejected_rows)}. Справочник не изменен."
        )
    if zero_cost:
        parts.append(
            f"Позиций с полной себестоимостью 0 ₽: {len(zero_cost)} — их доходность "
            "показывается как «Нет себестоимости»."
        )
    return " ".join(parts)


def _limited_lines(lines: list[str]) -> str:
    shown = lines[:_WARNING_LIST_LIMIT]
    text = "\n".join(f"• {line}" for line in shown)
    remainder = len(lines) - len(shown)
    return text + (f"\n…и еще {remainder}" if remainder else "")


class CostCatalogTabOZPriceAnalyzerApp(CatalogAndCategoryOZPriceAnalyzerApp):
    """v0.5.7: move the cost catalog to its own tab before Settings."""

    def __init__(self, *args, **kwargs) -> None:
        self._catalog_splitter_base_upper = 150
        # Rows of the last rejected catalog XLSX; kept only until restart.
        self.rejected_cost_catalog_rows: list[str] = []
        self.rejected_cost_catalog_source = ""
        super().__init__(*args, **kwargs)
        self._install_cost_catalog_tab()
        self.refresh_products()

    def _install_cost_catalog_tab(self) -> None:
        old_tree = self.products_tree
        self._hide_legacy_catalog_from_settings(old_tree)

        self.cost_catalog_tab = ttk.Frame(self.notebook, padding=4)
        settings_index = self.notebook.index(self.settings_tab)
        self.notebook.insert(
            settings_index,
            self.cost_catalog_tab,
            text="Справочник себестоимости",
        )

        self.cost_catalog_tab.columnconfigure(0, weight=1)
        self.cost_catalog_tab.rowconfigure(0, weight=1)
        self.cost_catalog_tab.rowconfigure(2, weight=0)

        top = ttk.Frame(self.cost_catalog_tab)
        top.grid(row=0, column=0, sticky="new", pady=(8, 0))
        top.columnconfigure(0, weight=1)

        header = ttk.Frame(top)
        header.grid(row=0, column=0, sticky="ew")
        header.columnconfigure(0, weight=1)
        ttk.Label(
            header,
            text="Справочник себестоимости",
            style="Section.TLabel",
        ).grid(row=0, column=0, sticky="w")
        ttk.Label(
            header,
            text="Текущие товары, категории и себестоимость для будущих расчетов",
            style="Muted.TLabel",
        ).grid(row=1, column=0, sticky="w", pady=(2, 0))

        actions = ttk.Frame(top)
        actions.grid(row=1, column=0, sticky="ew", pady=(10, 6))
        ttk.Button(
            actions,
            text="Редактировать справочник",
            style="Accent.TButton",
            command=self.open_cost_catalog_editor,
        ).grid(row=0, column=0, padx=(0, 4))
        ttk.Button(
            actions,
            text="Добавить товар",
            command=self.add_product,
        ).grid(row=0, column=1, padx=4)
        ttk.Button(
            actions,
            text="Изменить выбранный",
            command=self.edit_product,
        ).grid(row=0, column=2, padx=4)
        ttk.Button(
            actions,
            text="В архив / восстановить",
            command=self.toggle_product,
        ).grid(row=0, column=3, padx=4)
        ttk.Button(
            actions,
            text="Журнал изменений",
            command=self.show_cost_history,
        ).grid(row=0, column=4, padx=4)
        ttk.Button(
            actions,
            text="Удалить",
            style="Danger.TButton",
            command=self.delete_selected_product,
        ).grid(row=0, column=5, padx=(4, 0))

        self.catalog_help_label = ttk.Label(
            top,
            text=(
                "Справочник влияет только на будущие расчеты. "
                "Сохраненные отчеты хранят исторический снимок себестоимости."
            ),
            style="Muted.TLabel",
        )
        self.catalog_help_label.grid(row=2, column=0, sticky="w", pady=(0, 8))

        # Permanent warning; shown only while some rows have no full cost.
        self.cost_catalog_warning_frame = ttk.Frame(top)
        self.cost_catalog_warning_frame.grid(row=4, column=0, sticky="ew", pady=(0, 6))
        self.cost_catalog_warning_frame.columnconfigure(1, weight=1)
        self.cost_catalog_warning_var = tk.StringVar(master=self, value="")
        self.cost_catalog_warning_button = ttk.Button(
            self.cost_catalog_warning_frame,
            text=SHOW_SKIPPED_ROWS_TEXT,
            command=self.show_cost_catalog_warnings,
        )
        self.cost_catalog_warning_button.grid(row=0, column=0, sticky="w", padx=(0, 12))
        self.cost_catalog_warning_label = ttk.Label(
            self.cost_catalog_warning_frame,
            textvariable=self.cost_catalog_warning_var,
            style="Warning.TLabel",
            wraplength=760,
            justify="left",
        )
        self.cost_catalog_warning_label.grid(row=0, column=1, sticky="w")
        self.cost_catalog_warning_frame.grid_remove()

        filters = ttk.Frame(top)
        filters.grid(row=3, column=0, sticky="ew", pady=(0, 4))
        filters.columnconfigure(5, weight=1)

        ttk.Label(filters, text="Поиск:").grid(row=0, column=0, padx=(0, 6))
        search_entry = ttk.Entry(
            filters,
            textvariable=self.product_search_var,
            width=28,
        )
        search_entry.grid(row=0, column=1, padx=(0, 12))

        ttk.Label(filters, text="Показывать:").grid(row=0, column=2, padx=(0, 6))
        status_combo = ttk.Combobox(
            filters,
            textvariable=self.product_status_var,
            state="readonly",
            values=("Все", "Активные", "Архив"),
            width=12,
        )
        status_combo.grid(row=0, column=3, padx=(0, 12))
        status_combo.bind(
            "<<ComboboxSelected>>",
            lambda _event: self.refresh_products(),
        )

        ttk.Label(
            filters,
            textvariable=self.product_count_var,
            style="Muted.TLabel",
        ).grid(row=0, column=4, sticky="w")

        ttk.Button(
            filters,
            text="Выгрузить XLSX",
            command=self.export_product_catalog,
        ).grid(row=0, column=6, padx=4)
        ttk.Button(
            filters,
            text="Загрузить XLSX",
            command=self.import_product_catalog,
        ).grid(row=0, column=7, padx=4)
        ttk.Button(
            filters,
            text="Очистить справочник",
            command=self.clear_product_catalog,
        ).grid(row=0, column=8, padx=(12, 4))

        self.products_tree = self._create_tree(
            self.cost_catalog_tab,
            ["article", "name", "category", "total", "material", "labor", "status"],
            [
                "Артикул",
                "Наименование",
                "Категория",
                "Полная себестоимость",
                "Материал",
                "Трудозатраты",
                "Статус",
            ],
            row=2,
            widths=[150, 320, 220, 180, 150, 150, 110],
        )
        self.products_tree.bind("<Double-1>", lambda _event: self.edit_product())
        self.products_tree.bind(
            "<Delete>",
            lambda _event: self.delete_selected_product(),
            add="+",
        )

        def compact(upper: int) -> None:
            _hide_label_with_text(
                top,
                "Справочник влияет только на будущие расчеты",
                upper < 185,
            )

        self._catalog_splitter = _GridSplitter(
            self,
            self.cost_catalog_tab,
            row=1,
            table_row=2,
            absorb_row=0,
            min_upper=self._catalog_splitter_base_upper,
            min_table=220,
            compact=compact,
            protected_rows=(0,),
        )

        if hasattr(self, "_table_modes"):
            self._table_modes.pop("settings", None)
            self._table_modes["catalog"] = _TableModeController(
                self,
                "catalog",
                "Справочник себестоимости",
                self.cost_catalog_tab,
                self.products_tree,
                self._catalog_splitter,
                action_label="Изменить выбранный",
                action=self.edit_product,
            )

        factor = resolve_ui_scale(
            self._saved_scale_preference(),
            int(self.winfo_screenheight()),
        )
        self._catalog_splitter.min_upper = max(
            90,
            int(round(self._catalog_splitter_base_upper * factor)),
        )
        self._catalog_splitter._apply()

    def refresh_products(self) -> None:
        super().refresh_products()
        self._refresh_cost_catalog_warning()

    def _cost_catalog_warning_lines(self) -> tuple[list[str], list[str]]:
        zero_cost = catalog_positions_without_full_cost(self.db.list_products())
        return zero_cost, list(getattr(self, "rejected_cost_catalog_rows", []))

    def _refresh_cost_catalog_warning(self) -> None:
        frame = getattr(self, "cost_catalog_warning_frame", None)
        if frame is None:
            return
        zero_cost, rejected = self._cost_catalog_warning_lines()
        text = cost_catalog_warning_text(
            zero_cost,
            rejected,
            getattr(self, "rejected_cost_catalog_source", ""),
        )
        self.cost_catalog_warning_var.set(text)
        try:
            if text:
                frame.grid()
            else:
                frame.grid_remove()
        except tk.TclError:
            return
        guard = getattr(self, "_schedule_controls_guard", None)
        if callable(guard):
            guard()

    def _remember_rejected_cost_catalog(self, source: str) -> None:
        super()._remember_rejected_cost_catalog(source)
        self.rejected_cost_catalog_rows = rows_without_full_cost(source)
        self.rejected_cost_catalog_source = Path(source).name if self.rejected_cost_catalog_rows else ""
        self._refresh_cost_catalog_warning()

    def _forget_rejected_cost_catalog(self) -> None:
        super()._forget_rejected_cost_catalog()
        self.rejected_cost_catalog_rows = []
        self.rejected_cost_catalog_source = ""

    def show_cost_catalog_warnings(self) -> None:
        zero_cost, rejected = self._cost_catalog_warning_lines()
        if not zero_cost and not rejected:
            messagebox.showinfo(
                "Справочник себестоимости",
                "Строк без полной себестоимости нет.",
                parent=self,
            )
            return
        sections: list[str] = []
        if rejected:
            source = self.rejected_cost_catalog_source
            sections.append(
                f"XLSX «{source}» не загружен из-за строк без полной себестоимости:\n"
                + _limited_lines(rejected)
                + "\n\nЗаполните полную себестоимость и загрузите исправленный файл."
            )
        if zero_cost:
            sections.append(
                "Позиции справочника с полной себестоимостью 0 ₽:\n"
                + _limited_lines(zero_cost)
                + "\n\nИх доходность показывается как «Нет себестоимости». "
                "Укажите себестоимость кнопкой «Изменить выбранный»."
            )
        messagebox.showwarning(
            "Пропущенные строки справочника",
            "\n\n".join(sections),
            parent=self,
        )

    def _hide_legacy_catalog_from_settings(self, old_tree: ttk.Treeview) -> None:
        button = _button_with_text(self.settings_tab, "Редактировать справочник")
        if button is not None:
            try:
                button.master.grid_remove()
            except tk.TclError:
                pass

        try:
            old_tree.master.grid_remove()
        except tk.TclError:
            pass

        if hasattr(self, "_settings_splitter"):
            self._settings_splitter.suspended = True
            try:
                self._settings_splitter.bar.grid_remove()
            except tk.TclError:
                pass

        for row in (2, 3, 4):
            self.settings_tab.rowconfigure(row, weight=0, minsize=0)
        self.settings_tab.rowconfigure(1, weight=0, minsize=0)
        self.settings_tab.rowconfigure(5, weight=1, minsize=0)

    def _apply_ui_scale(self, factor: float) -> None:
        super()._apply_ui_scale(factor)
        if not hasattr(self, "_catalog_splitter"):
            return
        factor = min(1.0, max(0.8, float(factor)))
        self._catalog_splitter.min_upper = max(
            90,
            int(round(self._catalog_splitter_base_upper * factor)),
        )
        self._catalog_splitter._apply()


def run_app() -> None:
    app = CostCatalogTabOZPriceAnalyzerApp()
    app.mainloop()
